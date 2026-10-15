"""Excel export of what-if scenarios, for the client conversation."""

import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..models import Client

HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
BOLD = Font(bold=True)
INR = '[$₹-4009] #,##,##0'
PCT = "0.0%"


def _header(ws, row: int, cols: list[str]) -> None:
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font, cell.fill = HEAD, HEAD_FILL


def _widths(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def whatif_xlsx(client: Client, results: list[tuple[str, dict]], meta: dict, disclaimer: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = f"What-if review: {client.name}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Prepared {date.today():%d %b %Y} by {meta.get('distributor_name')}, " \
               f"AMFI-registered Mutual Fund Distributor, {meta.get('arn')}"
    ws["A3"] = disclaimer
    ws["A3"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A3:F3")
    ws.row_dimensions[3].height = 45

    r = 5
    _header(ws, r, ["Item"] + [name for name, _ in results])
    rows = [
        ("Gross sale value", "gross_proceeds", INR),
        ("Estimated exit loads", "exit_loads", INR),
        ("Estimated STT", "stt", INR),
        ("Estimated additional tax this FY", "incremental_tax", INR),
        ("Estimated amount after tax and costs", "net_proceeds", INR),
    ]
    for label, key, fmt in rows:
        r += 1
        ws.cell(row=r, column=1, value=label)
        for j, (_, res) in enumerate(results, 2):
            c = ws.cell(row=r, column=j, value=res[key])
            c.number_format = fmt
    r += 1
    ws.cell(row=r, column=1, value="Total tax for the FY after this change")
    for j, (_, res) in enumerate(results, 2):
        ws.cell(row=r, column=j, value=res["tax_after"]["total"]).number_format = INR
    r += 2
    ws.cell(row=r, column=1, value="Allocation after the change").font = BOLD
    classes = sorted({a["name"] for _, res in results for a in res["allocation_after"] + res["allocation_before"]})
    r += 1
    _header(ws, r, ["Asset class", "Now"] + [name for name, _ in results])
    before = {a["name"]: a["pct"] for a in results[0][1]["allocation_before"]} if results else {}
    for cls in classes:
        r += 1
        ws.cell(row=r, column=1, value=cls.title())
        ws.cell(row=r, column=2, value=before.get(cls, 0)).number_format = PCT
        for j, (_, res) in enumerate(results, 3):
            after = {a["name"]: a["pct"] for a in res["allocation_after"]}
            ws.cell(row=r, column=j, value=after.get(cls, 0)).number_format = PCT
    r += 2
    if results:
        a = results[0][1]["assumptions"]
        ws.cell(row=r, column=1, value="Assumptions").font = BOLD
        for k, v in (("Marginal slab rate", f"{a['slab_rate']:.0%}"), ("Surcharge", f"{a['surcharge_rate']:.0%}"),
                     ("Tax rules", a["rules"]), ("Prices as of", str(a["prices_asof"]))):
            r += 1
            ws.cell(row=r, column=1, value=k)
            ws.cell(row=r, column=2, value=v)
    _widths(ws, [42, 22, 22, 22, 22, 22])

    for name, res in results:
        sh = wb.create_sheet(name[:28])
        _header(sh, 1, ["Instrument", "Bought on", "Units", "Cost", "Sale value", "Gain", "Days held",
                        "Tax category", "Exit load"])
        rr = 1
        for line in res["lines"]:
            for lot in line["lots"]:
                rr += 1
                vals = [line["name"], lot["acquired"], lot["units"], lot["cost"], lot["proceeds"], lot["gain"],
                        lot["days_held"], lot["bucket_label"], lot["exit_load"]]
                for j, v in enumerate(vals, 1):
                    c = sh.cell(row=rr, column=j, value=v)
                    if j in (4, 5, 6, 9):
                        c.number_format = INR
        rr += 2
        sh.cell(row=rr, column=1, value="Tax by category (whole FY, after this change)").font = BOLD
        for b in res["tax_after"]["buckets"]:
            if b["gross"] or b["tax"]:
                rr += 1
                sh.cell(row=rr, column=1, value=b["label"])
                sh.cell(row=rr, column=4, value=b["gross"]).number_format = INR
                sh.cell(row=rr, column=6, value=b["tax"]).number_format = INR
        for w in res["warnings"]:
            rr += 1
            sh.cell(row=rr, column=1, value=f"Note: {w}")
        _widths(sh, [48, 12, 12, 14, 14, 14, 10, 40, 12])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
